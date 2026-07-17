import os
import time
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel, Field

load_dotenv()
my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("API Key is missing")

client = Groq(api_key=my_api_key)

model = "openai/gpt-oss-120b"

# part 1 => Extracting Data from the Job Description and storing it in the form of JSON

#JD
job_description = """
Description
Do you want to solve real customer problems through innovative technology? Do you enjoy working on scalable services in a collaborative team environment? Do you want to see your code directly impact millions of customers worldwide?

At Amazon, we hire the best minds in technology to innovate and build on behalf of our customers. Customer obsession is part of our company DNA, which has made us one of the world's most beloved brands.

Our Software Development Engineers (SDEs) use modern technology to solve complex problems while seeing their work's impact first-hand. The challenges SDEs solve at Amazon are meaningful and influence millions of customers, sellers, and products globally. We seek individuals passionate about creating new products, features, and services while managing ambiguity in an environment where development cycles are measured in weeks, not years.

At Amazon, we believe in ownership at every level. As an SDE-I, you'll own the entire lifecycle of your code - from design through deployment and ongoing operations. This ownership mindset, combined with our commitment to operational excellence, ensures we deliver the highest quality solutions for our customers.

We're looking for curious minds who think big and want to define tomorrow's technology. At Amazon, you'll grow into the high-impact engineer you know you can be, supported by a culture of learning and mentorship. Every day brings exciting new challenges and opportunities for personal growth.
Key job responsibilities
• Collaborate and communicate effectively with experienced cross-disciplinary Amazonians to design, build, and operate innovative products and services that delight our customers, while participating in technical discussions to drive solutions forward.
• Design and develop scalable solutions using cloud-native architectures and microservices in a large distributed computing environment.
• Participate in code reviews and contribute to technical documentation.
• Build and maintain resilient distributed systems that are scalable, fault-tolerant, and cost-effective.
• Leverage and contribute to the development of GenAI and AI-powered tools to enhance development productivity while staying current with emerging technologies.
• Write clean, maintainable code following best practices and design patterns.
• Work in an agile environment practicing CI/CD principles while participating in operational responsibilities including on-call duties.
• Demonstrate operational excellence through monitoring, troubleshooting, and resolving production issues.
Basic Qualifications
- Experience with at least one general-purpose programming language such as Java, Python, C++, C#, Go, Rust, or TypeScript
- Experience with data structure implementation, basic algorithm development, and/or object-oriented design principles
- Currently has, or is in the process of obtaining a bachelor’s degree in Computer Science, Computer Engineering, Data Science, Information Systems, or related STEM fields
- Must be 18 years of age of older
Preferred Qualifications
- Experience from previous technical internship(s) or demonstrated project experience
- Experience with one or more of the following: AI tools for development productivity, Cloud platforms (preferably AWS), Database systems (SQL and NoSQL), Contributing to open-source projects, Version control systems, Debugging and troubleshooting complex systems
- Demonstrated ability to learn and adapt to new technologies quickly
- Basic understanding of software development lifecycle (SDLC)
- Strong problem-solving and analytical skills
- Excellent written and verbal communication skills
 """

#job Schema Schema
class JobDescription(BaseModel):
    role : str
    required_skills : list[str]
    preferred_skills : list[str]
    min_exp : float | None
    max_exp : float | None
    education_requirements : list[str]
    responsibilities : list[str]

jobD_schema = JobDescription.model_json_schema()

system_prompt = f"""
You are an expert HR assistant.

Your job is to analyze job descriptions and extract
structured information from them.

Return ONLY valid JSON matching this schema:

{jobD_schema}
IMPORTANT:
Do NOT return the schema itself.
Do NOT return fields like "properties", "title" or "type".
Fill the schema with actual information extracted from the job description.

If minimum experience is not mentioned, return null.
If information for a list is missing, return an empty list.
Do not invent information.
"""

user_prompt =  f"""
Analyze the following job description:

{job_description}
"""
message_system = {
    "role" : "system",
    "content" : system_prompt
}

message_user = {
    "role" : "user",
    "content" : user_prompt
}

response_format = {
    "type" : "json_object"
}

messsages = [message_system, message_user]

response = client.chat.completions.create(model = model, messages = messsages, response_format = response_format)

answer = response.choices[0].message.content

raw_json = answer
# Isko phadte ksa hae
import json

job_data = json.loads(raw_json)
job = JobDescription(**job_data)

print(job.min_exp)
print(job.education_requirements)

# Part 1 => Job Description Parser Done

#Part 2 => Resume schema defining
class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company : str | None = None
    role: str | None = None
    duration:str | None = None
    description:str | None = None
    skills_used:list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email:str | None = None
    phone : int | None = None
    total_experience_years : float | None = None
    skills : list[str] = []
    experiences : list[Experience] = []
    education : list[str] = []
    projects : list[str] = []
    certifications : list[str] = []
    socials : list[str] = []

resume_schema = Resume.model_json_schema()

def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format, temperature = 0)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume

#part 2 done

#part 3 => extracting resume data

from pypdf import PdfReader
from docx import Document

def read_pfd(filepath) : 
    reader = PdfReader(filepath)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + '\n'
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text
    
def read_resume(filePath):
    if filePath.suffix.lower() == ".pdf":
        return read_pfd(filePath)
    elif filePath.suffix.lower() == ".docx":
        return read_docx(filePath)
    else:
        raise ValueError("Unsupported file format")    

# Part 3 Done

# part 4 => Parcing resume form the resume folder

resume_folder  =  Path("resumes")
all_results = []
for filePath in resume_folder.iterdir():
    if filePath.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\n Processing " + filePath.name + "")
    resume_text = read_resume(filePath)
    parsed_resume = parse_resume(resume_text)
    time.sleep(5)#to avoid API rate limiting for 1-2 resume is fine but for thousans or millians will jam the server basically network securitu from DOS and overcrouding => Rate Limiting
    result = final_score(job, parsed_resume)
    time.sleep(5)#to avoid API rate limiting
    print("Score", result.score)
    all_results.append({
        "name":parsed_resume.name,
        "score":result.score,
        "details":result.details
    })
    all_results.sort(key=lambda candidate: candidate["score"] , reverse=True)
    top2 = all_results[:2]
    worst2 = all_results[-2:]
    
    print("Top 2 candidates")
    for candidate in top2:
        print(f"{candidate['name']}\n{candidate['score']}%")
        print(candidate["details"])
    
    for candidate in worst2:
        print(f"{candidate['name']}\n{candidate['score']}%")
        print(candidate["details"])